import json
import io
import os
import random
from typing import List, Optional

import docx
import openpyxl
import PyPDF2
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse
from fastapi.encoders import jsonable_encoder
import os
# Import your custom modules
from model.manager.model_manager import DownloadResponse, HuggingFaceModel, ModelDownloadRequest, ModelManager
from utils.memory_extractor import extract_memory
from utils.remove_rag_tag import remove_specific_block
from utils.title_extractor import extract_title_block
from model.api_called import GeminiAgent
from prompts.chat_prompt import build_context
from services.library_services import RAGServices
from services.memory_services import MemoryServices
from utils.python_file import generate_text, generate_docx, generate_excel, generate_pdf
from link_services import get_all_urls_metadata
from utils.image_utils import load_base64_image
from global_storage import set_app_storage_dir, get_app_storage_dir
from Types.profile_type import EditProfileData, ProfileData
from files_services import check_if_userprofile_exists, load_userProfile_json, update_userProfile_json, create_userProfile_json
from services.chat_services import ChatServices

# Initialize services
model_manager = ModelManager()
gemini_api_key = os.environ.get("GEMINI_API_KEY", "")
app = FastAPI()

# CORS Middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Exception handler for validation errors
@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    print(f"Validation error on request: {request.url}")
    print(f"Details: {exc.errors()}")
    return JSONResponse(
        status_code=422,
        content=jsonable_encoder({"detail": exc.errors(), "body": exc.body}),
    )


# ================== HEALTH CHECK ROUTES ==================
@app.get("/health")
def health_check():
    """Health check endpoint"""
    random_number = random.randint(1, 100)
    return {"status": "success", "message": f"Here A random number {random_number}"}


@app.get("/ping")
def ping():
    """Simple ping endpoint"""
    return {"status": "ok"}


# ================== PROFILE MANAGEMENT ROUTES ==================
@app.post('/check_userprofile')
async def check_user_profile(request: Request):
    """Check and get profile data"""
    try:
        data = await request.json()
        path = data.get("path")
        
        if not path:
            return {"status": "failed", "message": "Missing path"}
        
        set_app_storage_dir(path)
        return check_if_userprofile_exists()
    
    except Exception as e:
        print(f"Error checking user profile: {e}")
        return {"status": "failed", "message": "Internal server error"}


@app.post("/create_profile")
def create_profile(profile: ProfileData):
    """Create a new profile file"""
    try:
        return create_userProfile_json(
            name=profile.name, 
            preferences=profile.preferences, 
            mode=profile.mode,
            image_path=profile.image_path if profile.image_path else None
        )
    except Exception as e:
        print(f"Error creating profile: {e}")
        return {"status": "failed", "message": "Failed to create profile"}


@app.put("/edit_profile")
async def edit_profile(request: Request):
    """Update profile content"""
    try:
        data = await request.json()
        new_data = data.get("new_data")
        
        if not new_data:
            return {"status": "failed", "message": "Missing new profile data"}
        
        return update_userProfile_json(new_data)
    
    except Exception as e:
        print(f"Error updating profile: {e}")
        return {"status": "failed", "message": "Failed to update profile"}


@app.get("/load_profile")
def load_profile():
    """Load profile data"""
    try:
        return load_userProfile_json()
    except Exception as e:
        print(f"Error loading profile: {e}")
        return {"status": "failed", "message": "Failed to load profile"}


@app.get("/load_profile_image")
def load_profile_image():
    """Load profile image"""
    try:
        profile_path = os.path.join(get_app_storage_dir(), "profile_image.png")
        image = load_base64_image(profile_path)
        
        if image:
            return {"status": "success", "message": image}
        else:
            return {"status": "failed", "message": "No profile image found"}
    
    except Exception as e:
        print(f"Error loading profile image: {e}")
        return {"status": "failed", "message": "Failed to load profile image"}


# ================== TOOLS ROUTES ==================
@app.post("/get_url_metadata")
async def get_url_metadata(request: Request):
    """Fetch URL metadata"""
    try:
        data = await request.json()
        urls = data.get("urls")
        
        if not urls:
            return {"status": "failed", "message": "Missing urls"}
        
        urls_metadata = get_all_urls_metadata(urls)
        return {"status": "success", "message": urls_metadata}
    
    except Exception as e:
        print(f"Error getting URL metadata: {e}")
        return {"status": "failed", "message": "Failed to fetch URL metadata"}


@app.post("/create_user_file")
async def create_user_file(request: Request):
    """Create user file"""
    try:
        data = await request.json()
        file_data = data.get("file_data")
        
        if not file_data:
            return {"status": "failed", "message": "Missing file data"}
        
        file_metadata = file_data.get("metadata", {})
        file_content = file_data.get("content", "")
        extension = file_metadata.get("extension", "").lower()
        file_name = file_metadata.get("file_name", "untitled")
        
        if extension == "txt":
            return generate_text(file_content, file_name)
        elif extension in ["doc", "docx"]:
            return generate_docx(file_content, file_name)
        elif extension == "xlsx":
            return generate_excel(file_content, file_name)
        elif extension == "pdf":
            return generate_pdf(file_content, file_name)
        else:
            return {"status": "failed", "message": "Unsupported file type"}
    
    except Exception as e:
        print(f"Error creating user file: {e}")
        return {"status": "failed", "message": "Failed to create file"}


# ================== CHAT MANAGEMENT ROUTES ==================
@app.post("/chat/new")
async def create_new_chat():
    """Create a new chat room"""
    try:
        chat_services = ChatServices()
        chat_id = chat_services.create_new_entry()
        
        if not chat_id:
            return {"status": "failed", "message": "Failed to create chat"}
        
        return {"status": "success", "message": chat_id}
    
    except Exception as e:
        print(f"Error creating new chat: {e}")
        return {"status": "failed", "message": "Failed to create chat"}


@app.put("/chat/rename")
async def rename_chat(chat_id: int, new_name: str):
    """Rename a chat room"""
    try:
        if chat_id <= 0:
            return {"status": "failed", "message": "Invalid chat ID"}
        
        if not new_name.strip():
            return {"status": "failed", "message": "New name cannot be empty"}
        
        chat_services = ChatServices()
        result = chat_services.update_chat_title(chat_id, new_title=new_name)
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error renaming chat: {e}")
        return {"status": "failed", "message": "Failed to rename chat"}


@app.get("/chat/load-all")
def load_all_chats():
    """Load all chat rooms"""
    try:
        chat_services = ChatServices()
        result = chat_services.load_chats_list()
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error loading chats: {e}")
        return {"status": "failed", "message": "Failed to load chats"}


@app.delete("/chat/delete-one")
def delete_chat(chat_id: int):
    """Delete a chat item"""
    try:
        if chat_id <= 0:
            return {"status": "failed", "message": "Invalid chat ID"}
        
        chat_services = ChatServices()
        result = chat_services.delete_chat_entry(chat_id)
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error deleting chat: {e}")
        return {"status": "failed", "message": "Failed to delete chat"}


@app.put("/chat/archive")
def toggle_chat_archive(current_archive_state: bool, chat_id: int):
    """Toggle archive state for chat"""
    try:
        if chat_id <= 0:
            return {"status": "failed", "message": "Invalid chat ID"}
        
        chat_services = ChatServices()
        result = chat_services.toggle_archive_chat(
            chat_id, 
            current_state=1 if current_archive_state else 0
        )
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error toggling chat archive: {e}")
        return {"status": "failed", "message": "Failed to toggle archive state"}


@app.get("/chat/search")
def search_chats(query: str):
    """Search for chats"""
    try:
        if not query.strip():
            return {"status": "failed", "message": "Search query cannot be empty"}
        
        chat_services = ChatServices()
        result = chat_services.search_for_chat(query)
        
        if not result:
            return {"status": "failed", "message": "Couldn't find any chats that align"}
        
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error searching chats: {e}")
        return {"status": "failed", "message": "Failed to search chats"}


@app.get("/chat/load-archived")
def load_archived_chats():
    """Load archived chat rooms"""
    try:
        chat_services = ChatServices()
        result = chat_services.load_all_archived_chats()
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error loading archived chats: {e}")
        return {"status": "failed", "message": "Failed to load archived chats"}


@app.get("/chat/save")
def save_chat(chat_id: int):
    """Save chat locally"""
    try:
        if chat_id <= 0:
            return {"status": "failed", "message": "Invalid chat ID"}
        
        chat_services = ChatServices()
        result = chat_services.save_chat_locally(chat_id)
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error saving chat: {e}")
        return {"status": "failed", "message": "Failed to save chat"}


# ================== MESSAGES ROUTES ==================
@app.post("/messages/new")
async def create_new_messages(chat_id: int, request: Request):
    """Save new message and model reply"""
    try:
        if chat_id <= 0:
            return {"status": "failed", "message": "Invalid chat ID"}
        
        data = await request.json()
        user_message = data.get("user_message")
        model_message = data.get("model_message")
        
        if not user_message or not model_message:
            return {"status": "failed", "message": "Missing message content"}
        
        chat_services = ChatServices()
        result = chat_services.create_new_message_entry(user_message, model_message, chat_id)
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error creating new messages: {e}")
        return {"status": "failed", "message": "Failed to create messages"}


@app.get("/messages/get-prev-context")
def get_previous_context(chat_id: int, original_message_id: int):
    """Get previous context for message regeneration"""
    try:
        if chat_id <= 0 or original_message_id <= 0:
            return {"status": "failed", "message": "Invalid ID provided"}
        
        chat_services = ChatServices()
        old_context = chat_services.get_context_for_regeneration(chat_id, original_message_id)
        
        if not old_context:
            return {"status": "failed", "message": "Couldn't fetch prior messages"}
        
        return {"status": "success", "message": old_context}
    
    except Exception as e:
        print(f"Error getting previous context: {e}")
        return {"status": "failed", "message": "Failed to fetch context"}


@app.post("/messages/regenerate")
async def regenerate_message_endpoint(chat_id: int, original_message_id: int, original_reply_id: int, request: Request):
    """Regenerate a message"""
    try:
        if chat_id <= 0 or original_message_id <= 0 or original_reply_id <= 0:
            return {"status": "failed", "message": "Invalid ID provided"}
        
        data = await request.json()
        user_message = data.get("user_message")
        model_message = data.get("model_message")
        
        if not user_message or not model_message:
            return {"status": "failed", "message": "Missing message content"}
        
        chat_services = ChatServices()
        updated_messages = chat_services.regenerate_message(
            chat_id, user_message, model_message, original_message_id, original_reply_id
        )
        
        if not updated_messages:
            return {"status": "failed", "message": "Couldn't regenerate message"}
        
        return {"status": "success", "message": updated_messages}
    
    except Exception as e:
        print(f"Error regenerating message: {e}")
        return {"status": "failed", "message": "Failed to regenerate message"}


@app.get("/messages/load-all")
def load_all_messages(chat_id: int):
    """Load all messages for a chat"""
    try:
        if chat_id <= 0:
            return {"status": "failed", "message": "Invalid chat ID"}
        
        chat_services = ChatServices()
        result = chat_services.load_all_chat_messages(chat_id)
        
        if not result:
            return {"status": "failed", "message": "Couldn't load messages"}
        
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error loading messages: {e}")
        return {"status": "failed", "message": "Failed to load messages"}


@app.get("/messages/search")
def search_messages(query: str):
    """Search for messages"""
    try:
        if not query.strip():
            return {"status": "failed", "message": "Search query cannot be empty"}
        
        chat_services = ChatServices()
        result = chat_services.search_for_message(query)
        
        if not result:
            return {"status": "failed", "message": "Couldn't find any messages that meet condition"}
        
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error searching messages: {e}")
        return {"status": "failed", "message": "Failed to search messages"}


@app.put('/messages/archive')
def toggle_message_archive(current_archive_state: bool, message_id: int):
    """Toggle archive state for a message"""
    try:
        if message_id <= 0:
            return {"status": "failed", "message": "Invalid message ID"}
        
        chat_services = ChatServices()
        result = chat_services.toggle_archive_message(
            message_id, 
            current_state=1 if current_archive_state else 0
        )
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error toggling message archive: {e}")
        return {"status": "failed", "message": "Failed to toggle archive state"}


@app.get("/messages/load-archived")
def load_archived_messages():
    """Load archived messages"""
    try:
        chat_services = ChatServices()
        result = chat_services.load_all_archived_messages()
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error loading archived messages: {e}")
        return {"status": "failed", "message": "Failed to load archived messages"}


# ================== MEDIA ROUTES ==================
@app.get("/media/load-all")
async def load_all_media():
    """Load all media from database"""
    try:
        chat_services = ChatServices()
        data = chat_services.fetch_media_content()
        return {"status": "success", "message": data}
    
    except Exception as e:
        print(f"Error loading media: {e}")
        return {"status": "failed", "message": "Failed to load media"}


# ================== AI MODEL ENGINE ROUTES ==================
@app.post("/model/chat/new")
async def chat_with_model(
    chat_id: int, 
    need_title: bool, 
    useRag: bool, 
    request: Request, 
    mode: Optional[str] = None, 
    action: Optional[str] = None
):
    """Create new chat message with AI model"""
    try:
        print("Starting chat with model...")
        
        data = await request.json()
        user_message = data.get("user_message")
        user_settings = data.get("user_settings", {})
        
        if not user_message:
            return {"status": "failed", "message": "Missing user message"}
        
        # Initialize services
        chat_services = ChatServices()
        rag_services = RAGServices()
        memory_services = MemoryServices()
        
        old_context = []
        chat_item = None
        
        # Get context if not a temporary chat
        if chat_id != -1:
            old_context = chat_services.load_n_chat_messages(chat_id, 6)
        
        # Create new chat if it's a temporary chat
        if chat_id == -1:
            chat_item = chat_services.create_new_entry(title="New Chat")
            if not chat_item:
                return {"status": "failed", "message": "Failed to create chat"}
        
        # Build context and get model response
        print("Building context...")
        prompt = build_context(
            chat_history=old_context,
            memory_service=memory_services,
            ragServices=rag_services,
            current_user_input=user_message,
            need_title=need_title,
            include_rag=useRag,
            mode=mode,
            user_settings=user_settings,
            chat_id=chat_id
        )
        
        # Get model response
        print("Getting model response...")
        agent = GeminiAgent(
            api_key=gemini_api_key,
            system_prompt=prompt
        )
        response = await agent.run(user_message)
        
        # Process memory extraction
        parsed_response = extract_memory(response["final"])
        if parsed_response["memory_item"]:
            memory_services.save(parsed_response["memory_item"], chat_id)
        
        current_chat_id = chat_item[0]["id"] if (chat_item and len(chat_item) > 0) else chat_id
        
        # Update chat title if needed
        if need_title:
            title = extract_title_block(parsed_response["message_item"])
            if title:
                chat_services.update_chat_title(current_chat_id, title)
        
        # Save messages
        if current_chat_id == -1:
            return {"status": "failed", "message": "Chat id supplied is out of reasonable range (-1)"}
        
        result = chat_services.create_new_message_entry(
            user_message, parsed_response["message_item"], current_chat_id
        )
        
        return {"status": "success", "message": {"response": result, "chat": chat_item}}
    
    except Exception as e:
        print(f"Error in chat with model: {e}")
        return {"status": "failed", "message": "Failed to process chat request"}


@app.post("/model/chat/update")
async def update_message_with_model(chat_id: int, original_message_id: int, original_reply_id: int, request: Request):
    """Update message with AI model"""
    try:
        if chat_id <= 0 or original_message_id <= 0 or original_reply_id <= 0:
            return {"status": "failed", "message": "Invalid ID provided"}
        
        data = await request.json()
        updated_message = data.get("updated_message")
        
        if not updated_message:
            return {"status": "failed", "message": "Missing updated message"}
        
        # Get context and services
        chat_services = ChatServices()
        rag_services = RAGServices()
        old_context = chat_services.get_context_for_regeneration(chat_id, original_message_id)
        
        # Build prompt
        prompt = build_context(
            chat_history=old_context,
            memory_service=MemoryServices(),
            library_service=rag_services,
            current_user_input=updated_message,
            chat_id=chat_id
        )
        
        # Get model response
        agent = GeminiAgent(
            api_key=gemini_api_key,
            system_prompt=prompt
        )
        new_response = await agent.run(updated_message)
        
        # Save updated message
        result = chat_services.regenerate_message(
            chat_id, updated_message, new_response["final"], original_message_id, original_reply_id
        )
        
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error updating message: {e}")
        return {"status": "failed", "message": "Failed to update message"}


@app.post("/model/chat/re")
async def regenerate_with_model(
    chat_id: int, 
    original_message_id: int, 
    original_reply_id: int, 
    use_rag: bool, 
    request: Request,
    mode: Optional[str] = None, 
    action: Optional[str] = None
):
    """Regenerate message with AI model"""
    try:
        if chat_id <= 0 or original_message_id <= 0 or original_reply_id <= 0:
            return {"status": "failed", "message": "Invalid ID provided"}
        
        data = await request.json()
        user_message = data.get("user_message")
        model_reply = data.get("model_replay")  # Note: keeping original typo for compatibility
        user_settings = data.get("user_settings", {})
        
        if not user_message or not model_reply:
            return {"status": "failed", "message": "Missing message content"}
        
        create_input = f"user:{user_message} assistant:{model_reply} answer the user message in a different style and way try"
        
        # Get context and services
        chat_services = ChatServices()
        rag_services = RAGServices()
        memory_services = MemoryServices()
        old_context = chat_services.get_context_for_regeneration(chat_id, original_message_id)
        
        # Build prompt
        prompt = build_context(
            chat_history=old_context,
            memory_service=memory_services,
            ragServices=rag_services,
            current_user_input=create_input,
            user_settings=user_settings,
            need_title=False,
            mode=mode,
            include_rag=use_rag,
            chat_id=chat_id
        )
        
        # Get model response
        agent = GeminiAgent(
            api_key=gemini_api_key,
            system_prompt=prompt
        )
        new_response = await agent.run(create_input)
        
        # Save regenerated message
        result = chat_services.regenerate_message(
            chat_id, create_input, new_response["final"], original_message_id, original_reply_id
        )
        
        return {"status": "success", "message": result}
    
    except Exception as e:
        print(f"Error regenerating message: {e}")
        return {"status": "failed", "message": "Failed to regenerate message"}


# ================== RAG (Retrieval-Augmented Generation) ROUTES ==================
@app.post("/rag/upload")
async def upload_rag_file(metadata: str = Form(...), file: UploadFile = File(...)):
    """Upload file for RAG processing"""
    try:
        # Parse metadata
        try:
            meta = json.loads(metadata)
        except json.JSONDecodeError:
            return {"status": "failed", "message": "Invalid metadata format"}
        
        extension = meta.get("extension", "").lower()
        if not extension:
            return {"status": "failed", "message": "Missing file extension"}
        
        # Read file content
        raw_bytes = await file.read()
        content = ""
        
        # Process different file types
        if extension == "txt":
            content = raw_bytes.decode("utf-8")
        
        elif extension == "json":
            try:
                parsed = json.loads(raw_bytes.decode("utf-8"))
                content = json.dumps(parsed, indent=2)
            except json.JSONDecodeError:
                return {"status": "failed", "message": "Invalid JSON file"}
        
        elif extension == "docx":
            try:
                doc = docx.Document(io.BytesIO(raw_bytes))
                content = "\n".join([p.text for p in doc.paragraphs])
            except Exception:
                return {"status": "failed", "message": "Failed to read DOCX file"}
        
        elif extension == "pdf":
            try:
                reader = PyPDF2.PdfReader(io.BytesIO(raw_bytes))
                content = ""
                for page in reader.pages:
                    content += page.extract_text() or ""
            except Exception:
                return {"status": "failed", "message": "Failed to extract PDF text"}
        
        elif extension == "xlsx":
            try:
                wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True)
                content = ""
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        content += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
            except Exception:
                return {"status": "failed", "message": "Failed to read Excel file"}
        
        else:
            return {"status": "failed", "message": f"Unsupported file type: {extension}"}
        
        # Save to database
        rag_services = RAGServices()
        rag_services.save_to_db(
            file_text=content,
            filename=meta.get("filename", file.filename or "unknown"),
            extension=extension,
            title=meta.get("title", ""),
            is_isolated=meta.get("is_isolated", False),
            chat_id=str(meta.get("chat_id", "")),
            tags=meta.get("tags", "")
        )
        
        return {"status": "success"}
    
    except Exception as e:
        print(f"Error uploading RAG file: {e}")
        return {"status": "failed", "message": "Failed to upload file"}


@app.get("/rag/files")
def list_rag_files():
    """List all RAG files"""
    try:
        rag_services = RAGServices()
        return rag_services.load_all_rag_files()
    except Exception as e:
        print(f"Error listing RAG files: {e}")
        return {"status": "failed", "message": "Failed to load RAG files"}


@app.get("/rag/search")
def search_rag(query: str):
    """Search RAG database"""
    try:
        if not query.strip():
            return {"status": "failed", "message": "Search query cannot be empty"}
        
        rag_services = RAGServices()
        return rag_services.rag_query(query)
    except Exception as e:
        print(f"Error searching RAG: {e}")
        return {"status": "failed", "message": "Failed to search RAG database"}


@app.delete("/rag/files/delete/{file_id}")
def delete_rag_file(file_id: int):
    """Delete a RAG file"""
    try:
        if file_id <= 0:
            return {"status": "failed", "message": "Invalid file ID"}
        
        rag_services = RAGServices()
        rag_services.remove_file(file_id)
        return {"status": "deleted"}
    except Exception as e:
        print(f"Error deleting RAG file: {e}")
        return {"status": "failed", "message": "Failed to delete file"}


# ================== MEMORY MANAGEMENT ROUTES ==================
@app.get("/memories/all")
def load_all_memories():
    """Load all memories"""
    try:
        memory_services = MemoryServices()
        memories = memory_services.get_all()
        return {"status": "success", "message": memories}
    except Exception as e:
        print(f"Error loading memories: {e}")
        return {"status": "failed", "message": "Failed to load memories"}


@app.put("/memories/update/{memory_id}")
async def update_memory(memory_id: int, request: Request):
    """Update existing memory"""
    try:
        if memory_id <= 0:
            return {"status": "failed", "message": "Invalid memory ID"}
        
        data = await request.json()
        new_memory_content = data.get("updated_content")
        
        if not new_memory_content:
            return {"status": "failed", "message": "Missing updated content"}
        
        memory_services = MemoryServices()
        memory_services.update(memory_id, new_memory_content)
        return {"status": "success", "message": "updated"}
    except Exception as e:
        print(f"Error updating memory: {e}")
        return {"status": "failed", "message": "Failed to update memory"}


@app.delete("/memories/delete/{memory_id}")
def delete_memory(memory_id: int):
    """Delete a memory"""
    try:
        if memory_id <= 0:
            return {"status": "failed", "message": "Invalid memory ID"}
        
        memory_services = MemoryServices()
        memory_services.delete(memory_id)
        return {"status": "success", "message": "deleted"}
    except Exception as e:
        print(f"Error deleting memory: {e}")
        return {"status": "failed", "message": "Failed to delete memory"}


@app.post("/memories/add")
async def add_new_memory(request: Request):
    """Add new memory manually"""
    try:
        data = await request.json()
        new_memory_content = data.get("memory_content")
        new_memory_weight = data.get("memory_weight", 1)  # Default weight
        
        if not new_memory_content:
            return {"status": "failed", "message": "Missing memory content"}
        
        memory_services = MemoryServices()
        item = memory_services.create_memory_manually(new_memory_content, new_memory_weight)
        return {"status": "success", "message": item}
    except Exception as e:
        print(f"Error adding memory: {e}")
        return {"status": "failed", "message": "Failed to add memory"}


# ================== LOCAL MODELS MANAGEMENT ROUTES ==================
@app.get("/api/models/huggingface", response_model=List[HuggingFaceModel])
async def get_huggingface_models(
    search: Optional[str] = None,
    task: Optional[str] = "text-generation", 
    limit: int = 20,
    sort: str = "downloads"
):
    """Get list of models from HuggingFace Hub"""
    try:
        if limit <= 0 or limit > 100:
            limit = 20  # Set reasonable default
        
        return await model_manager.get_huggingface_models(search, task, limit, sort)
    except Exception as e:
        print(f"Error getting HuggingFace models: {e}")
        raise HTTPException(status_code=500, detail="Failed to fetch HuggingFace models")


@app.post("/api/models/download-local", response_model=DownloadResponse)
async def download_local_model(request: ModelDownloadRequest, background_tasks: BackgroundTasks):
    """Download a popular local model"""
    try:
        if not request.modelId:
            raise HTTPException(status_code=400, detail="Model ID is required")
        
        return await model_manager.download_local_model(request.modelId)
    except Exception as e:
        print(f"Error downloading local model: {e}")
        raise HTTPException(status_code=500, detail="Failed to download local model")


@app.post("/api/models/download-huggingface", response_model=DownloadResponse)
async def download_huggingface_model(request: ModelDownloadRequest):
    """Download a model from HuggingFace (for GGUF files)"""
    try:
        if not request.modelId:
            raise HTTPException(status_code=400, detail="Model ID is required")
        
        return await model_manager.download_huggingface_model(request.modelId)
    except Exception as e:
        print(f"Error downloading HuggingFace model: {e}")
        raise HTTPException(status_code=500, detail="Failed to download HuggingFace model")


@app.get("/api/models/local")
async def list_local_models():
    """List locally downloaded models"""
    try:
        return await model_manager.list_local_models()
    except Exception as e:
        print(f"Error listing local models: {e}")
        raise HTTPException(status_code=500, detail="Failed to list local models")


@app.delete("/api/models/local/{model_id}")
async def delete_local_model(model_id: str):
    """Delete a locally downloaded model"""
    try:
        if not model_id.strip():
            raise HTTPException(status_code=400, detail="Model ID is required")
        
        return await model_manager.delete_local_model(model_id)
    except Exception as e:
        print(f"Error deleting local model: {e}")
        raise HTTPException(status_code=500, detail="Failed to delete local model")


@app.get("/api/models/popular")
async def get_popular_models():
    """Get list of popular models available for download"""
    try:
        return await model_manager.get_popular_models()
    except Exception as e:
        print(f"Error getting popular models: {e}")
        raise HTTPException(status_code=500, detail="Failed to fetch popular models")

